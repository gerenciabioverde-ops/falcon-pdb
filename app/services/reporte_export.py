"""Descarga del REPORTE (relación de lo que se presenta a SUNAT) en Excel y PDF."""
from __future__ import annotations
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models import ESTADOS, Declaracion
from app.services import archivos_sunat as arch
from app.services import pdb as svc

MESES = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Setiembre", "Octubre", "Noviembre", "Diciembre"]

COL_VENTAS = ["N°", "Fecha", "Tipo", "Comprobante", "Cliente", "Documento", "Moneda", "Venta", "Base S/", "IGV S/", "Total S/", "Estado"]
COL_COMPRAS = ["N°", "Fecha", "Tipo", "Comprobante", "Proveedor", "RUC", "Mon. orig.", "Base S/", "IGV S/", "Otros S/", "Total S/", "Detracción"]
COL_PAGOS = ["N°", "Fecha de operación", "Tipo", "Comprobante pagado", "RUC / Doc.", "Medio de pago", "Banco", "N° de operación", "Monto S/"]


def _filas(rep: dict) -> dict[str, list[list]]:
    d = lambda x: float(x)  # noqa: E731
    ventas = [[i, v["fecha"], v["tipo"], v["doc"], v["nombre"], v["ruc"], v["moneda"], "Exportación" if v["externa"] else "Interna",
               d(v["base"]), d(v["igv"]), d(v["total"]), "ANULADO" if v["anulado"] else "Válido"]
              for i, v in enumerate(rep["lista_ventas"], 1)]
    compras = [[i, c["fecha"], c["tipo"], c["doc"], c["nombre"], c["ruc"], c["moneda"], d(c["base"]), d(c["igv"]), d(c["otros"]), d(c["total"]),
                c["detraccion"]] for i, c in enumerate(rep["lista_compras"], 1)]
    pagos = [[i, p["fecha"], p["tipo"], p["doc"], p["ruc"], p["medio"], p["banco"], p["operacion"], d(p["monto"])]
             for i, p in enumerate(rep["lista_pagos"], 1)]
    return {"Ventas": ventas, "Compras": compras, "Pagos": pagos}


def _resumen(decl: Declaracion, rep: dict) -> list[list]:
    c, v, du, e = rep["compras"], rep["ventas"], rep["duas"], rep["cdps"]
    f = lambda x: float(x)  # noqa: E731
    filas = [["Compras", c["n"], f(c["base"]), f(c["igv"]), f(c["total"])],
             ["Forma de pago", rep["pagos"]["n"], None, None, f(rep["pagos"]["monto"])],
             ["Ventas", v["n"], f(v["base"]), f(v["igv"]), f(v["total"])],
             ["DUAS · bloque D (US$ FOB)", du["n"], None, None, f(du["fob"])],
             ["DUAS · bloque E (US$ FOB)", e["n"], None, None, f(e["fob"])],
             ["Tipo de cambio (días)", rep["tc"], None, None, None]]
    return filas


def _titulo(decl: Declaracion) -> str:
    return f"Reporte PDB Exportadores {decl.periodo} · {decl.empresa.razon_social} · RUC {decl.empresa.ruc}"


def a_excel(decl: Declaracion) -> bytes:
    rep = svc.reporte(decl)
    wb = Workbook()
    verde = PatternFill("solid", fgColor="065F46")
    claro = PatternFill("solid", fgColor="D1FAE5")
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = _titulo(decl)
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Período {decl.periodo} ({MESES[decl.mes]} {decl.anio}) · Estado: {ESTADOS[decl.estado]}"
    ws.append([])
    ws.append(["Hoja", "Filas", "Base imponible S/", "IGV S/", "Total"])
    for c in ws[4]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), verde
    for fila in _resumen(decl, rep):
        ws.append(fila)
    ws.append([])
    ws.append(["Archivos que se presentan", "Líneas"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for nombre, lineas in arch.archivos(decl).items():
        ws.append([nombre, len(lineas)])
    ws.append([])
    ws.append(["Observaciones de la validación"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for h in svc.validar(decl):
        ws.append([("ERROR: " if h["nivel"] == "error" else "Aviso: ") + h["texto"]])
    for fila in ws.iter_rows(min_row=5, max_row=10, min_col=3, max_col=5):
        for c in fila:
            c.number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 38
    for col in "BCDE":
        ws.column_dimensions[col].width = 20

    datos = _filas(rep)
    for nombre, cols in (("Ventas", COL_VENTAS), ("Compras", COL_COMPRAS), ("Pagos", COL_PAGOS)):
        hoja = wb.create_sheet(nombre)
        hoja.append(cols)
        for c in hoja[1]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), verde, Alignment(horizontal="center", wrap_text=True)
        for fila in datos[nombre]:
            hoja.append(fila)
        n = len(datos[nombre])
        numericas = [i for i, c in enumerate(cols, 1) if c.endswith("S/")]
        if n:
            tot = ["TOTAL"] + [""] * (len(cols) - 1)
            for i in numericas:
                L = get_column_letter(i)
                tot[i - 1] = f"=SUM({L}2:{L}{n + 1})"
            hoja.append(tot)
            for c in hoja[n + 2]:
                c.font, c.fill = Font(bold=True), claro
        for i in numericas:
            for r in range(2, n + 3):
                hoja.cell(r, i).number_format = "#,##0.00"
        for i, c in enumerate(cols, 1):
            ancho = max([len(str(c))] + [len(str(f[i - 1])) for f in datos[nombre]]) + 2
            hoja.column_dimensions[get_column_letter(i)].width = min(max(ancho, 8), 42)
        hoja.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _m(x) -> str:
    return "" if x in (None, "") else f"{Decimal(str(x)):,.2f}"


def a_pdf(decl: Declaracion) -> bytes:
    rep = svc.reporte(decl)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title=_titulo(decl))
    ss = getSampleStyleSheet()
    chico = ParagraphStyle("chico", parent=ss["BodyText"], fontSize=7, leading=8.5)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=11, spaceBefore=8, spaceAfter=3)
    verde = colors.HexColor("#065F46")
    estilo = lambda n_tot=False: TableStyle([  # noqa: E731
        ("FONTSIZE", (0, 0), (-1, -1), 7), ("BACKGROUND", (0, 0), (-1, 0), verde), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        *([("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#D1FAE5")), ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold")] if n_tot else []),
    ])
    P = lambda t: Paragraph(str(t).replace("&", "&amp;").replace("<", "&lt;"), chico)  # noqa: E731
    el = [Paragraph(f"<b>Reporte de la información a presentar a SUNAT · PDB Exportadores</b>", ss["Title"]),
          Paragraph(f"{decl.empresa.razon_social} · RUC {decl.empresa.ruc} · Período <b>{decl.periodo}</b> ({MESES[decl.mes]} {decl.anio}) · "
                    f"Estado: {ESTADOS[decl.estado]}", ss["BodyText"])]
    el.append(Paragraph("1. Resumen por hoja", h2))
    t = [["Hoja", "Filas", "Base imponible S/", "IGV S/", "Total"]] + [[a, b, _m(c), _m(d), _m(e)] for a, b, c, d, e in _resumen(decl, rep)]
    tb = Table(t, colWidths=[70 * mm, 20 * mm, 40 * mm, 35 * mm, 40 * mm])
    tb.setStyle(estilo())
    tb.setStyle(TableStyle([("ALIGN", (1, 0), (-1, -1), "RIGHT")]))
    el.append(tb)
    el.append(Paragraph("2. Archivos que se presentan", h2))
    t = [["Archivo", "Líneas"]] + [[n, len(l)] for n, l in arch.archivos(decl).items()] or [["(sin datos)", ""]]
    tb = Table(t, colWidths=[90 * mm, 25 * mm])
    tb.setStyle(estilo())
    el.append(tb)

    datos = _filas(rep)
    def tabla(titulo, cols, filas, anchos, num_idx, total_fila):
        el.append(Paragraph(titulo, h2))
        t = [cols] + [[P(c) if isinstance(c, str) and len(c) > 14 else (_m(c) if i in num_idx and isinstance(c, float) else c)
                       for i, c in enumerate(f)] for f in filas]
        if filas:
            t.append(total_fila)
        tb = Table(t, colWidths=[a * mm for a in anchos], repeatRows=1)
        tb.setStyle(estilo(bool(filas)))
        tb.setStyle(TableStyle([("ALIGN", (i, 0), (i, -1), "RIGHT") for i in num_idx]))
        el.append(tb)
    c, v = rep["compras"], rep["ventas"]
    tabla(f"3. Relación de comprobantes de VENTAS ({len(datos['Ventas'])})", COL_VENTAS, datos["Ventas"],
          [8, 17, 9, 28, 55, 22, 12, 20, 25, 18, 25, 20], {8, 9, 10},
          ["TOTAL", "", "", "", "", "", "", "", _m(v["base"]), _m(v["igv"]), _m(v["total"]), ""])
    tabla(f"4. Relación de comprobantes de COMPRAS ({len(datos['Compras'])})", COL_COMPRAS, datos["Compras"],
          [8, 17, 9, 28, 55, 22, 14, 22, 18, 16, 24, 25], {7, 8, 9, 10},
          ["TOTAL", "", "", "", "", "", "", _m(c["base"]), _m(c["igv"]), _m(c["otros"]), _m(c["total"]), ""])
    tabla(f"5. Relación de PAGOS ({len(datos['Pagos'])})", COL_PAGOS, datos["Pagos"],
          [8, 28, 9, 40, 28, 25, 18, 40, 25], {8},
          ["TOTAL", "", "", "", "", "", "", "", _m(rep["pagos"]["monto"])])
    el.append(Paragraph("6. Observaciones de la validación", h2))
    for h in svc.validar(decl) or [{"nivel": "ok", "texto": "Sin observaciones."}]:
        el.append(Paragraph(("✖ " if h["nivel"] == "error" else ("▲ " if h["nivel"] == "aviso" else "✔ ")) + str(h["texto"]).replace("&", "&amp;"), chico))
    doc.build(el)
    return buf.getvalue()
